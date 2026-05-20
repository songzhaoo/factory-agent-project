from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import logging
import os
import re
import urllib.request
import uuid
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any
from urllib.parse import quote, urlencode

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:  # Docker installs openpyxl; this keeps local syntax checks usable.
    load_workbook = None

from app.db import Database, now_text
from app.schemas import (
    AgentAskResponse,
    DashboardSummary,
    Equipment,
    ImportResult,
    InjectionRun,
    InventoryItem,
    Mold,
    MultimodalRecord,
    Order,
    ProductionTask,
    QualityIssue,
    TaskDetail,
    WorkReportRecord,
    WorkReportResponse,
    WorkReportSubmit,
)

logger = logging.getLogger(__name__)

MULTIMODAL_CATEGORIES = {
    "document": "单据识别",
    "quality": "质检图片分析",
    "equipment": "设备现场识别",
    "inventory": "库存标签盘点",
    "drawing": "图纸/工艺单理解",
}

PROCESS_KEYWORDS = [
    "深孔钻",
    "车床",
    "铣床",
    "磨床",
    "摇臂钻床",
    "CNC加工",
    "火花机",
    "线切割",
    "钳工装配",
    "试模",
    "注塑生产",
    "质检",
    "包装出货",
]

SHIPPER_CODES = {
    "顺丰": "SF",
    "中通": "ZTO",
    "圆通": "YTO",
    "申通": "STO",
    "韵达": "YD",
    "极兔": "JTSD",
    "京东": "JD",
    "德邦": "DBL",
    "邮政": "YZPY",
    "EMS": "EMS",
}

PROVINCE_NAMES = {
    "河北", "山西", "辽宁", "吉林", "黑龙江", "江苏", "浙江", "安徽", "福建", "江西", "山东", "河南",
    "湖北", "湖南", "广东", "海南", "四川", "贵州", "云南", "陕西", "甘肃", "青海", "台湾",
    "内蒙古", "广西", "西藏", "宁夏", "新疆", "北京", "天津", "上海", "重庆", "香港", "澳门",
}


def rows_to_models(rows: list[dict[str, Any]], model: type) -> list:
    return [model(**dict(row)) for row in rows]


class FactoryDataService:
    def __init__(self, database: Database) -> None:
        self.database = database
        self.frontend_base_url = os.getenv("FRONTEND_BASE_URL", "http://127.0.0.1:8900").rstrip("/")
        self.upload_dir = Path(os.getenv("UPLOAD_DIR", "data/uploads"))
        self.upload_dir.mkdir(parents=True, exist_ok=True)

    def dashboard(self) -> DashboardSummary:
        with self.database.connect() as conn:
            order_count = conn.execute("SELECT COUNT(*) c FROM orders").fetchone()["c"]
            task_count = conn.execute("SELECT COUNT(*) c FROM production_tasks").fetchone()["c"]
            running_task_count = conn.execute("SELECT COUNT(*) c FROM production_tasks WHERE current_status = '加工中'").fetchone()["c"]
            exception_count = conn.execute("SELECT COUNT(*) c FROM quality_issues WHERE status != '已关闭'").fetchone()["c"]
            low_inventory_count = conn.execute("SELECT COUNT(*) c FROM inventory_items WHERE qty < safety_qty").fetchone()["c"]
            overdue_risk_count = conn.execute("SELECT COUNT(*) c FROM production_tasks WHERE current_status != '已完成' AND planned_end_at < ?", (now_text(),)).fetchone()["c"]
        return DashboardSummary(
            order_count=order_count,
            task_count=task_count,
            running_task_count=running_task_count,
            exception_count=exception_count,
            low_inventory_count=low_inventory_count,
            overdue_risk_count=overdue_risk_count,
        )

    def list_orders(self) -> list[Order]:
        with self.database.connect() as conn:
            rows = conn.execute("SELECT * FROM orders ORDER BY updated_at DESC").fetchall()
        return [self._to_order(row) for row in rows]

    def list_molds(self) -> list[Mold]:
        return self._list("molds", Mold)

    def list_equipment(self) -> list[Equipment]:
        return self._list("equipment", Equipment)

    def list_tasks(self, order_id: str | None = None) -> list[ProductionTask]:
        with self.database.connect() as conn:
            if order_id:
                rows = conn.execute(
                    "SELECT * FROM production_tasks WHERE order_id = ? ORDER BY process_seq, updated_at DESC",
                    (order_id,),
                ).fetchall()
            else:
                rows = conn.execute("SELECT * FROM production_tasks ORDER BY process_seq, updated_at DESC").fetchall()
        return [self._to_task(row) for row in rows]

    def list_injection_runs(self) -> list[InjectionRun]:
        return self._list("injection_runs", InjectionRun)

    def list_quality_issues(self) -> list[QualityIssue]:
        return self._list("quality_issues", QualityIssue)

    def list_inventory(self) -> list[InventoryItem]:
        return self._list("inventory_items", InventoryItem)

    def list_multimodal_records(self) -> list[MultimodalRecord]:
        with self.database.connect() as conn:
            rows = conn.execute("SELECT * FROM multimodal_records ORDER BY created_at DESC LIMIT 50").fetchall()
        return [self._to_multimodal_record(row) for row in rows]

    def create_multimodal_record(self, category: str, filename: str, content_type: str | None, content: bytes, description: str = "") -> MultimodalRecord:
        if category not in MULTIMODAL_CATEGORIES:
            raise ValueError("不支持的多模态类型")
        if not content:
            raise ValueError("上传文件不能为空")
        if len(content) > 8 * 1024 * 1024:
            raise ValueError("图片不能超过 8MB")
        ext = Path(filename).suffix.lower()
        if ext not in {".jpg", ".jpeg", ".png", ".webp", ".bmp"}:
            raise ValueError("仅支持 jpg、jpeg、png、webp、bmp 图片")
        record_id = uuid.uuid4().hex
        safe_name = re.sub(r"[^0-9A-Za-z_.-]", "_", filename)[:120] or f"{record_id}{ext}"
        storage_path = self.upload_dir / f"{record_id}_{safe_name}"
        storage_path.write_bytes(content)
        analysis = self._analyze_multimodal(category, filename, description, content, content_type)
        now = now_text()
        with self.database.connect() as conn:
            conn.execute(
                """
                INSERT INTO multimodal_records
                (id, category, filename, content_type, storage_path, file_size, description,
                 extracted_json, analysis, suggestions, source, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    record_id,
                    category,
                    filename,
                    content_type,
                    str(storage_path),
                    len(content),
                    description,
                    json.dumps(analysis["extracted"], ensure_ascii=False),
                    analysis["analysis"],
                    analysis["suggestions"],
                    analysis["source"],
                    now,
                ),
            )
            row = conn.execute("SELECT * FROM multimodal_records WHERE id = ?", (record_id,)).fetchone()
        logger.info("multimodal_record_created id=%s category=%s filename=%s size=%s", record_id, category, filename, len(content))
        return self._to_multimodal_record(row)

    def get_task_detail(self, task_id: str) -> TaskDetail | None:
        with self.database.connect() as conn:
            task = conn.execute("SELECT * FROM production_tasks WHERE id = ?", (task_id,)).fetchone()
            if task is None:
                return None
            reports = conn.execute(
                "SELECT * FROM work_reports WHERE task_id = ? ORDER BY created_at DESC LIMIT 30",
                (task_id,),
            ).fetchall()
        return TaskDetail(task=self._to_task(task), reports=rows_to_models(reports, WorkReportRecord))

    def submit_report(self, task_id: str, payload: WorkReportSubmit) -> WorkReportResponse | None:
        status_after = {
            "开始加工": "加工中",
            "加工完成": "已完成",
            "异常暂停": "异常暂停",
        }[payload.action]
        now = now_text()
        with self.database.connect() as conn:
            task = conn.execute("SELECT * FROM production_tasks WHERE id = ?", (task_id,)).fetchone()
            if task is None:
                logger.warning("work_report_task_not_found task_id=%s action=%s operator=%s", task_id, payload.action, payload.operator_name)
                return None
            actual_start_at = task["actual_start_at"]
            actual_finish_at = task["actual_finish_at"]
            if payload.action == "开始加工" and not actual_start_at:
                actual_start_at = now
            if payload.action == "加工完成":
                actual_finish_at = now
            conn.execute(
                """
                UPDATE production_tasks
                SET current_status = ?, actual_start_at = ?, actual_finish_at = ?, remark = ?, updated_at = ?
                WHERE id = ?
                """,
                (status_after, actual_start_at, actual_finish_at, payload.remark, now, task_id),
            )
            conn.execute("UPDATE equipment SET status = ?, operator_name = ?, updated_at = ? WHERE code = ?", (status_after, payload.operator_name, now, task["equipment_code"]))
            if task["mold_id"]:
                conn.execute("UPDATE molds SET current_stage = ?, status = ?, owner_name = ?, updated_at = ? WHERE id = ?", (task["process_name"], status_after, payload.operator_name, now, task["mold_id"]))
            report_id = uuid.uuid4().hex
            conn.execute(
                "INSERT INTO work_reports VALUES (?, ?, ?, ?, ?, ?, ?)",
                (report_id, task_id, payload.action, status_after, payload.operator_name, payload.remark, now),
            )
            updated_task = conn.execute("SELECT * FROM production_tasks WHERE id = ?", (task_id,)).fetchone()
            report = conn.execute("SELECT * FROM work_reports WHERE id = ?", (report_id,)).fetchone()
        logger.info(
            "work_report_submitted task_id=%s action=%s status_after=%s operator=%s report_id=%s",
            task_id,
            payload.action,
            status_after,
            payload.operator_name,
            report_id,
        )
        return WorkReportResponse(
            message=f"已提交：{payload.action}",
            task=self._to_task(updated_task),
            report=WorkReportRecord(**dict(report)),
        )

    def import_table(self, data_type: str, filename: str, content: bytes) -> ImportResult:
        logger.info("import_started data_type=%s filename=%s size_bytes=%s", data_type, filename, len(content))
        rows = self._read_rows(filename, content)
        errors: list[str] = []
        success = 0
        with self.database.connect() as conn:
            for index, row in enumerate(rows, start=2):
                try:
                    if data_type == "orders":
                        self._upsert_order(conn, row)
                    elif data_type == "tasks":
                        self._upsert_task(conn, row)
                    elif data_type == "inventory":
                        self._upsert_inventory(conn, row)
                    elif data_type == "quality":
                        self._upsert_quality(conn, row)
                    else:
                        raise ValueError("不支持的导入类型")
                    success += 1
                except Exception as exc:
                    errors.append(f"第{index}行：{exc}")
                    logger.warning("import_row_failed data_type=%s filename=%s row=%s error=%s", data_type, filename, index, exc)
            batch_id = uuid.uuid4().hex
            conn.execute(
                "INSERT INTO import_batches VALUES (?, ?, ?, ?, ?, ?, ?)",
                (batch_id, data_type, filename, len(rows), success, len(errors), now_text()),
            )
        logger.info(
            "import_finished batch_id=%s data_type=%s filename=%s total_rows=%s success_rows=%s error_rows=%s",
            batch_id,
            data_type,
            filename,
            len(rows),
            success,
            len(errors),
        )
        return ImportResult(batch_id=batch_id, data_type=data_type, total_rows=len(rows), success_rows=success, error_rows=len(errors), errors=errors[:20])

    def _list(self, table: str, model: type) -> list:
        with self.database.connect() as conn:
            rows = conn.execute(f"SELECT * FROM {table} ORDER BY updated_at DESC" if table != "equipment" else f"SELECT * FROM {table} ORDER BY code").fetchall()
        return rows_to_models(rows, model)

    @staticmethod
    def _to_multimodal_record(row: dict[str, Any]) -> MultimodalRecord:
        data = dict(row)
        data["extracted"] = json.loads(data.pop("extracted_json") or "{}")
        return MultimodalRecord(**data)

    def _to_task(self, row: dict[str, Any]) -> ProductionTask:
        task = ProductionTask(**dict(row))
        task.report_url = f"{self.frontend_base_url}/work-report?task_id={task.id}"
        return task

    def _to_order(self, row: dict[str, Any]) -> Order:
        order = Order(**dict(row))
        order.report_url = f"{self.frontend_base_url}/work-report?order_id={order.id}"
        return order

    @staticmethod
    def _read_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
        if filename.lower().endswith(".csv"):
            text = content.decode("utf-8-sig")
            return [dict(row) for row in csv.DictReader(io.StringIO(text))]
        if load_workbook is None:
            raise RuntimeError("当前环境缺少 openpyxl，无法导入 .xlsx，请安装依赖或使用 CSV")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(item or "").strip() for item in values[0]]
        return [{headers[i]: value for i, value in enumerate(row) if i < len(headers) and headers[i]} for row in values[1:]]

    @staticmethod
    def _value(row: dict[str, Any], *names: str, default: Any = "") -> Any:
        for name in names:
            value = row.get(name)
            if value not in (None, ""):
                return value
        return default

    def _upsert_order(self, conn: Any, row: dict[str, Any]) -> None:
        now = now_text()
        order_no = str(self._value(row, "订单编号", "order_no"))
        if not order_no:
            raise ValueError("订单编号不能为空")
        values = (
            f"ORD-{order_no}",
            order_no,
            str(self._value(row, "产品名称", "product_name")),
            str(self._value(row, "产品类型", "product_type", default="塑胶件")),
            int(self._value(row, "数量", "quantity", default=0)),
            str(self._value(row, "交期", "due_date", default=now)),
            str(self._value(row, "状态", "status", default="未开始")),
            str(self._value(row, "优先级", "priority", default="普通")),
            now,
            now,
        )
        conn.execute(
            """
            INSERT INTO orders (id, order_no, product_name, product_type, quantity, due_date, status, priority, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON DUPLICATE KEY UPDATE product_name=VALUES(product_name), product_type=VALUES(product_type),
            quantity=VALUES(quantity), due_date=VALUES(due_date), status=VALUES(status), priority=VALUES(priority), updated_at=VALUES(updated_at)
            """,
            values,
        )

    def _upsert_task(self, conn: Any, row: dict[str, Any]) -> None:
        now = now_text()
        task_id = str(self._value(row, "任务编号", "id", default=uuid.uuid4().hex[:8]))
        values = (
            task_id,
            str(self._value(row, "产品", "产品名称", "product_name")),
            str(self._value(row, "工序", "process_name")),
            int(self._value(row, "工序顺序", "process_seq", default=99)),
            str(self._value(row, "设备", "equipment_code")),
            str(self._value(row, "状态", "current_status", default="未开始")),
            str(self._value(row, "负责人", "owner_name", default="待分配")),
            str(self._value(row, "计划开始", "planned_start_at", default="")),
            str(self._value(row, "计划结束", "planned_end_at", default="")),
            str(self._value(row, "备注", "remark", default="")),
            now,
            now,
        )
        conn.execute(
            """
            INSERT INTO production_tasks
            (id, product_name, process_name, process_seq, equipment_code, current_status, owner_name, planned_start_at, planned_end_at, remark, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON DUPLICATE KEY UPDATE product_name=VALUES(product_name), process_name=VALUES(process_name),
            process_seq=VALUES(process_seq), equipment_code=VALUES(equipment_code), current_status=VALUES(current_status),
            owner_name=VALUES(owner_name), planned_start_at=VALUES(planned_start_at), planned_end_at=VALUES(planned_end_at),
            remark=VALUES(remark), updated_at=VALUES(updated_at)
            """,
            values,
        )

    def _upsert_inventory(self, conn: Any, row: dict[str, Any]) -> None:
        now = now_text()
        material = str(self._value(row, "材料名称", "material_name"))
        if not material:
            raise ValueError("材料名称不能为空")
        conn.execute(
            """
            INSERT INTO inventory_items (id, material_name, spec, unit, qty, safety_qty, supplier, eta, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                uuid.uuid4().hex,
                material,
                str(self._value(row, "规格", "spec", default="")),
                str(self._value(row, "单位", "unit", default="kg")),
                float(self._value(row, "库存", "qty", default=0)),
                float(self._value(row, "安全库存", "safety_qty", default=0)),
                str(self._value(row, "供应商", "supplier", default="")),
                str(self._value(row, "预计到货", "eta", default="")),
                now,
            ),
        )

    def _upsert_quality(self, conn: Any, row: dict[str, Any]) -> None:
        now = now_text()
        conn.execute(
            """
            INSERT INTO quality_issues
            (id, product_name, issue_type, bad_qty, cause, status, owner_name, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                uuid.uuid4().hex,
                str(self._value(row, "产品", "产品名称", "product_name")),
                str(self._value(row, "异常类型", "issue_type")),
                int(self._value(row, "不良数量", "bad_qty", default=0)),
                str(self._value(row, "原因", "cause", default="")),
                str(self._value(row, "状态", "status", default="处理中")),
                str(self._value(row, "负责人", "owner_name", default="")),
                now,
                now,
            ),
        )

    @staticmethod
    def _analyze_multimodal(category: str, filename: str, description: str, content: bytes, content_type: str | None) -> dict[str, Any]:
        vision_result = FactoryDataService._analyze_with_dashscope_vision(category, filename, description, content, content_type)
        if vision_result:
            return vision_result
        category_name = MULTIMODAL_CATEGORIES[category]
        text = f"{filename} {description}".strip()
        extracted: dict[str, Any] = {
            "资料类型": category_name,
            "文件名": filename,
            "备注": description,
        }
        if category == "document":
            extracted.update(
                {
                    "订单号": FactoryDataService._find_token(text, r"(SO\d{6,}|\bORD[-_A-Za-z0-9]+\b)") or "待人工确认",
                    "单据类型": FactoryDataService._match_first(text, ["采购单", "送货单", "质检单", "报工单"]) or "待识别单据",
                    "数量": FactoryDataService._find_token(text, r"(\d+(?:\.\d+)?)(?:件|个|kg|KG|块)") or "待人工确认",
                    "供应商": FactoryDataService._find_token(text, r"([\u4e00-\u9fa5A-Za-z0-9]{2,20})(?:供应商|材料商)") or "待人工确认",
                }
            )
            result = "已完成单据图片归档，并提取订单、数量、供应商等候选字段。"
            suggestions = "建议人工复核关键字段后同步到订单、库存、质检或报工模块。"
        elif category == "quality":
            issue = FactoryDataService._match_first(text, ["缩水", "毛边", "划痕", "色差", "变形", "缺料", "烧焦"]) or "疑似外观异常"
            extracted.update({"缺陷类型": issue, "关联产品": FactoryDataService._guess_product(text), "严重程度": "待质检复核"})
            result = f"已归档质检图片，初步识别为{issue}，可关联质检异常记录。"
            suggestions = "建议质检员复核缺陷类型、补充不良数量，并检查注塑参数或模具状态。"
        elif category == "equipment":
            extracted.update(
                {
                    "设备编号": FactoryDataService._find_token(text.upper(), r"(CNC-\d+|EDM-\d+|WIRE-\d+|INJ-\d+|DH-\d+)") or "待人工确认",
                    "报警代码": FactoryDataService._find_token(text.upper(), r"(ALARM[-_\s]?\d+|E\d{2,5})") or "待人工确认",
                    "运行状态": FactoryDataService._match_first(text, ["报警", "停机", "运行", "待机"]) or "待人工确认",
                }
            )
            result = "已归档设备现场图片，并提取设备编号、报警代码和运行状态候选字段。"
            suggestions = "建议维修人员复核报警代码，必要时调整设备排产并通知相关工序负责人。"
        elif category == "inventory":
            extracted.update(
                {
                    "物料名称": FactoryDataService._match_first(text.upper(), ["ABS", "PP", "P20", "PC", "PA", "POM"]) or "待人工确认",
                    "规格": FactoryDataService._find_token(text, r"(\d+\*\d+\*\d+|\d+(?:\.\d+)?\s?KG|\d+(?:\.\d+)?kg)") or "待人工确认",
                    "批次": FactoryDataService._find_token(text.upper(), r"(LOT[-_A-Z0-9]+|BATCH[-_A-Z0-9]+)") or "待人工确认",
                }
            )
            result = "已归档库存标签图片，并提取物料、规格和批次候选字段。"
            suggestions = "建议仓管复核数量后生成入库、出库或盘点记录。"
        else:
            extracted.update(
                {
                    "材料": FactoryDataService._match_first(text.upper(), ["P20", "718", "NAK80", "ABS", "PP"]) or "待人工确认",
                    "工序": FactoryDataService._match_first(text, PROCESS_KEYWORDS) or "待人工确认",
                    "尺寸/公差": FactoryDataService._find_token(text, r"(\d+(?:\.\d+)?\s?mm|[+-]0\.\d+)") or "待人工确认",
                }
            )
            result = "已归档图纸/工艺单图片，并提取材料、工序和尺寸公差候选字段。"
            suggestions = "建议工艺人员复核后生成工序任务或加工注意事项。"
        return {"extracted": extracted, "analysis": result, "suggestions": suggestions, "source": "local multimodal rules"}

    @staticmethod
    def _analyze_with_dashscope_vision(category: str, filename: str, description: str, content: bytes, content_type: str | None) -> dict[str, Any] | None:
        api_key = os.getenv("DASHSCOPE_API_KEY", "")
        if not api_key:
            return None
        model = os.getenv("DASHSCOPE_VL_MODEL", "qwen-vl-plus")
        mime = content_type or "image/png"
        image_data = base64.b64encode(content).decode("ascii")
        prompt = (
            f"你是工厂生产协同系统的多模态识别工具。图片类型：{MULTIMODAL_CATEGORIES[category]}。"
            f"文件名：{filename}。用户备注：{description or '无'}。"
            "请只输出JSON，格式为："
            "{\"extracted\":{\"字段\":\"值\"},\"analysis\":\"一句话识别结论\",\"suggestions\":\"处理建议\"}。"
            "字段需要尽量提取订单号、物料、数量、供应商、设备编号、报警代码、缺陷类型、工序、尺寸或公差等与图片相关的信息。"
        )
        body = {
            "model": model,
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": prompt},
                        {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{image_data}"}},
                    ],
                }
            ],
        }
        request = urllib.request.Request(
            "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions",
            data=json.dumps(body).encode("utf-8"),
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            method="POST",
        )
        try:
            with urllib.request.urlopen(request, timeout=30) as response:
                payload = json.loads(response.read().decode("utf-8"))
            content_text = payload["choices"][0]["message"]["content"]
            match = re.search(r"\{.*\}", content_text, flags=re.S)
            parsed = json.loads(match.group(0) if match else content_text)
            return {
                "extracted": parsed.get("extracted") or {},
                "analysis": str(parsed.get("analysis") or "已完成多模态图片识别。"),
                "suggestions": str(parsed.get("suggestions") or "建议人工复核关键字段后入库。"),
                "source": f"DashScope {model}",
            }
        except Exception as exc:
            logger.warning("dashscope_vision_failed model=%s category=%s filename=%s error=%s", model, category, filename, exc)
            return None

    @staticmethod
    def _find_token(text: str, pattern: str) -> str:
        match = re.search(pattern, text)
        return match.group(1).strip() if match else ""

    @staticmethod
    def _match_first(text: str, values: list[str]) -> str:
        upper_text = text.upper()
        for value in values:
            if value.upper() in upper_text:
                return value
        return ""

    @staticmethod
    def _guess_product(text: str) -> str:
        for value in ["洗衣机面板", "冰箱配件", "汽车配件", "水壶外壳", "茶杯"]:
            if value in text:
                return value
        return "待人工确认"


@dataclass
class AgentTool:
    name: str
    description: str
    runner: Callable[[dict[str, Any]], list[dict[str, Any]]]


class LLMToolRouter:
    def __init__(self) -> None:
        self.api_key = os.getenv("DASHSCOPE_API_KEY", "")
        self.model = os.getenv("DASHSCOPE_MODEL", "qwen-turbo")

    def choose_tool(self, question: str, tools: list[AgentTool]) -> str | None:
        if not self.api_key:
            return None
        tool_defs = [
            {
                "type": "function",
                "function": {
                    "name": tool.name,
                    "description": tool.description,
                    "parameters": {"type": "object", "properties": {}, "additionalProperties": True},
                },
            }
            for tool in tools
        ]
        body = {
            "model": self.model,
            "messages": [
                {"role": "system", "content": "你是工厂生产协同Agent，只选择最合适的工具，不要回答正文。"},
                {"role": "user", "content": question},
            ],
            "tools": tool_defs,
            "tool_choice": "auto",
        }
        request = urllib.request.Request(
            "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions",
            data=json.dumps(body).encode("utf-8"),
            headers={"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"},
            method="POST",
        )
        try:
            with urllib.request.urlopen(request, timeout=15) as response:
                payload = json.loads(response.read().decode("utf-8"))
            calls = payload["choices"][0]["message"].get("tool_calls") or []
            if calls:
                return calls[0]["function"]["name"]
        except Exception as exc:
            logger.warning("llm_tool_router_failed model=%s error=%s", self.model, exc)
            return None
        return None


class FactoryAgentService:
    def __init__(self, data_service: FactoryDataService) -> None:
        self.data_service = data_service
        self.llm_router = LLMToolRouter()
        self.tools = [
            AgentTool("query_order_progress", "查询订单和模具整体进度", self.query_order_progress),
            AgentTool("query_process_tasks", "查询深孔钻、车床、铣床、磨床、摇臂钻床、CNC、火花机、线切割、注塑等工序任务", self.query_process_tasks),
            AgentTool("query_equipment_schedule", "查询设备排产和设备当前任务", self.query_equipment_schedule),
            AgentTool("query_injection_runs", "查询注塑机生产进度", self.query_injection_runs),
            AgentTool("query_quality_issues", "查询质检异常和不良记录", self.query_quality_issues),
            AgentTool("query_inventory", "查询原料和模具钢库存", self.query_inventory),
            AgentTool("query_weather", "调用外部天气API查询城市天气，用于判断出货、送货和外发加工风险", self.query_weather),
            AgentTool("query_logistics", "调用快递鸟接口查询快递或物流单号的运输轨迹、签收状态和异常节点", self.query_logistics),
            AgentTool("query_route_plan", "调用高德地图Web服务查询发货地到收货地的驾车路线、距离和预计运输时长", self.query_route_plan),
            AgentTool("query_web_search", "调用Tavily联网搜索公开资料、行业信息、材料行情和外部知识", self.query_web_search),
            AgentTool("query_workday_calendar", "调用节假日API判断日期是否节假日、周末或工作日，用于交期风险判断", self.query_workday_calendar),
            AgentTool("query_multimodal_records", "查询多模态图片识别记录，包括单据、质检缺陷图、设备报警图、库存标签和工艺图纸", self.query_multimodal_records),
            AgentTool("generate_delay_risk_report", "生成延期风险报告", self.generate_delay_risk_report),
        ]

    def ask(self, question: str, context: dict[str, Any] | None = None) -> AgentAskResponse:
        active_context = self._normalize_context(context or {})
        effective_question = self._with_context(question, active_context)
        selected = self._force_tool_by_rules(question) or self.llm_router.choose_tool(effective_question, self.tools) or self._choose_tool_by_rules(effective_question)
        tool = next(item for item in self.tools if item.name == selected)
        data = tool.runner({"question": effective_question, "raw_question": question, "context": active_context})
        answer = self._format_answer(selected, data)
        next_context = self._context_from_result(question, data, active_context)
        self._log_tool(question, selected, {"context": active_context, "effective_question": effective_question})
        logger.info(
            "agent_answered tool=%s used_llm=%s rows=%s question_length=%s",
            selected,
            bool(self.llm_router.api_key),
            len(data),
            len(question),
        )
        return AgentAskResponse(answer=answer, tool=selected, data=data, context=next_context, used_llm=bool(self.llm_router.api_key))

    def query_order_progress(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        rows = [item.model_dump() for item in self.data_service.list_orders()] + [item.model_dump() for item in self.data_service.list_molds()]
        return self._filter_by_context(rows, args)

    def query_process_tasks(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        q = str(args.get("raw_question") or args.get("question") or "").upper()
        rows = [task.model_dump() for task in self.data_service.list_tasks()]
        keywords = ["深孔钻", "车床", "铣床", "磨床", "摇臂", "CNC", "火花", "线切割", "注塑", "洗衣机", "冰箱", "汽车", "水壶"]
        hits = [key for key in keywords if key.upper() in q]
        if hits:
            rows = [row for row in rows if any(key.upper() in f"{row['process_name']} {row['equipment_code']} {row['product_name']}".upper() for key in hits)]
        rows = self._filter_by_context(rows, args)
        return rows

    def query_equipment_schedule(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        return self.query_process_tasks(args)

    def query_injection_runs(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        return self._filter_by_context([item.model_dump() for item in self.data_service.list_injection_runs()], args)

    def query_quality_issues(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        return self._filter_by_context([item.model_dump() for item in self.data_service.list_quality_issues()], args)

    def query_inventory(self, _: dict[str, Any]) -> list[dict[str, Any]]:
        return [item.model_dump() for item in self.data_service.list_inventory()]

    def query_multimodal_records(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        q = str(args.get("raw_question") or args.get("question") or "")
        rows = [item.model_dump() for item in self.data_service.list_multimodal_records()]
        if any(key in q for key in ["单据", "送货单", "采购单", "报工单"]):
            return [row for row in rows if row["category"] == "document"] or rows
        if any(key in q for key in ["缺陷", "图片", "质检", "缩水", "毛边", "划痕", "色差"]):
            return [row for row in rows if row["category"] == "quality"] or rows
        if any(key in q for key in ["设备", "报警", "面板", "铭牌"]):
            return [row for row in rows if row["category"] == "equipment"] or rows
        if any(key in q for key in ["库存", "标签", "盘点", "物料"]):
            return [row for row in rows if row["category"] == "inventory"] or rows
        if any(key in q for key in ["图纸", "工艺", "尺寸", "公差"]):
            return [row for row in rows if row["category"] == "drawing"] or rows
        return rows

    def query_weather(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        question = str(args.get("question", ""))
        city = self._guess_city(question)
        # Open-Meteo is used as a no-key external API demo. It can be replaced with QWeather/HeWeather by key.
        geo_url = f"https://geocoding-api.open-meteo.com/v1/search?name={quote(city)}&count=1&language=zh&format=json"
        geo = self._load_external_json(geo_url)
        results = geo.get("results") or []
        if not results:
            return [{"city": city, "error": "没有查到城市经纬度"}]
        location = results[0]
        lat = location["latitude"]
        lon = location["longitude"]
        weather_url = (
            "https://api.open-meteo.com/v1/forecast"
            f"?latitude={lat}&longitude={lon}"
            "&daily=weather_code,temperature_2m_max,temperature_2m_min,precipitation_probability_max"
            "&forecast_days=3&timezone=Asia%2FShanghai"
        )
        weather = self._load_external_json(weather_url)
        daily = weather.get("daily") or {}
        rows = []
        for index, day in enumerate(daily.get("time", [])):
            rows.append(
                {
                    "city": location.get("name", city),
                    "date": day,
                    "weather_code": daily.get("weather_code", [None])[index],
                    "temp_max": daily.get("temperature_2m_max", [None])[index],
                    "temp_min": daily.get("temperature_2m_min", [None])[index],
                    "rain_probability": daily.get("precipitation_probability_max", [None])[index],
                    "source": "Open-Meteo external API",
                }
            )
        return rows

    def query_logistics(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        question = str(args.get("question", ""))
        ebusiness_id = os.getenv("KDNIAO_EBUSINESS_ID", "")
        api_key = os.getenv("KDNIAO_API_KEY", "")
        if not ebusiness_id or not api_key:
            return [{"error": "快递鸟接口未配置，请先配置 KDNIAO_EBUSINESS_ID 和 KDNIAO_API_KEY。"}]
        logistic_code = self._extract_logistic_code(question)
        shipper_code = self._guess_shipper_code(question)
        if not logistic_code:
            return [{"error": "没有识别到物流单号。请这样问：查询顺丰 1234567890 的物流状态。"}]
        if not shipper_code:
            return [{"error": "没有识别到快递公司。请补充顺丰、中通、圆通、申通、韵达、极兔、京东、德邦、邮政或 EMS。", "logistic_code": logistic_code}]

        request_data = json.dumps(
            {"OrderCode": "", "ShipperCode": shipper_code, "LogisticCode": logistic_code},
            ensure_ascii=False,
            separators=(",", ":"),
        )
        md5_text = hashlib.md5((request_data + api_key).encode("utf-8")).hexdigest()
        data_sign = base64.b64encode(md5_text.encode("utf-8")).decode("utf-8")
        form = {
            "RequestType": os.getenv("KDNIAO_REQUEST_TYPE", "8002"),
            "EBusinessID": ebusiness_id,
            "RequestData": request_data,
            "DataSign": data_sign,
            "DataType": "2",
        }
        try:
            payload = self._post_form_json("https://api.kdniao.com/Ebusiness/EbusinessOrderHandle.aspx", form)
        except Exception as exc:
            return [{"error": f"快递鸟接口调用失败：{exc}"}]
        traces = payload.get("Traces") or []
        return [
            {
                "logistic_code": logistic_code,
                "shipper_code": shipper_code,
                "success": payload.get("Success"),
                "state": payload.get("State"),
                "reason": payload.get("Reason") or payload.get("ReasonDesc") or "",
                "latest_accept_time": traces[-1].get("AcceptTime") if traces else "",
                "latest_accept_station": traces[-1].get("AcceptStation") if traces else "",
                "trace_count": len(traces),
                "traces": traces[-5:],
                "source": "KDNiao external API",
            }
        ]

    def query_route_plan(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        question = str(args.get("question", ""))
        api_key = os.getenv("AMAP_KEY", "")
        if not api_key:
            return [{"error": "高德地图接口未配置，请先配置 AMAP_KEY。"}]
        origin, destination = self._extract_route_points(question)
        if not origin or not destination:
            return [{"error": "没有识别到完整起点和终点。请这样问：从深圳工厂到东莞客户的送货路线多久？"}]
        if self._is_ambiguous_province(destination):
            return [{"error": f"目的地“{destination}”只到省/直辖市级别，路线距离会不稳定。请补充城市或详细地址，例如：从深圳到江西南昌送货多久。"}]
        try:
            origin_geo = self._amap_geocode(origin, api_key)
            destination_geo = self._amap_geocode(destination, api_key)
        except Exception as exc:
            return [{"origin": origin, "destination": destination, "error": f"高德地理编码调用失败：{exc}"}]
        if origin_geo.get("error") or destination_geo.get("error"):
            return [{"origin": origin, "destination": destination, "error": origin_geo.get("error") or destination_geo.get("error")}]
        route_url = (
            "https://restapi.amap.com/v3/direction/driving?"
            + urlencode({"key": api_key, "origin": origin_geo["location"], "destination": destination_geo["location"], "extensions": "base"})
        )
        try:
            route = self._load_external_json(route_url)
        except Exception as exc:
            return [{"origin": origin, "destination": destination, "error": f"高德路线规划调用失败：{exc}"}]
        if route.get("status") != "1":
            return [{"origin": origin, "destination": destination, "error": route.get("info") or "高德路线规划失败"}]
        paths = ((route.get("route") or {}).get("paths") or [])
        if not paths:
            return [{"origin": origin, "destination": destination, "error": "没有查询到可用路线"}]
        path = paths[0]
        distance_m = float(path.get("distance") or 0)
        duration_s = float(path.get("duration") or 0)
        return [
            {
                "origin": origin,
                "destination": destination,
                "origin_location": origin_geo["location"],
                "destination_location": destination_geo["location"],
                "distance_km": round(distance_m / 1000, 1),
                "duration_minutes": round(duration_s / 60),
                "strategy": path.get("strategy") or "默认驾车路线",
                "source": "AMap Web Service external API",
            }
        ]

    def query_web_search(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        question = str(args.get("raw_question") or args.get("question") or "")
        api_key = os.getenv("TAVILY_API_KEY", "")
        if not api_key:
            return [{"error": "Tavily 搜索接口未配置，请先配置 TAVILY_API_KEY。"}]
        try:
            payload = self._post_json(
                "https://api.tavily.com/search",
                {
                    "query": question,
                    "search_depth": "basic",
                    "max_results": 3,
                    "include_answer": True,
                    "include_raw_content": False,
                },
                {"Authorization": f"Bearer {api_key}"},
            )
        except Exception as exc:
            return [{"error": f"Tavily 搜索接口调用失败：{exc}"}]
        results = payload.get("results") or []
        return [
            {
                "query": question,
                "answer": payload.get("answer") or "",
                "results": [
                    {
                        "title": item.get("title") or "",
                        "url": item.get("url") or "",
                        "content": item.get("content") or "",
                        "score": item.get("score"),
                    }
                    for item in results[:3]
                ],
                "source": "Tavily external API",
            }
        ]

    def query_workday_calendar(self, args: dict[str, Any]) -> list[dict[str, Any]]:
        question = str(args.get("question", ""))
        api_key = os.getenv("JIEJIARI_API_KEY", "")
        if not api_key:
            return [{"error": "节假日接口未配置，请先配置 JIEJIARI_API_KEY。"}]
        target_date = self._extract_date(question)
        url = "https://api.jiejiariapi.com/v1/is_holiday?" + urlencode({"date": target_date, "key": api_key})
        try:
            payload = self._load_external_json(url)
        except Exception as exc:
            return [{"error": f"节假日接口调用失败：{exc}"}]
        return [
            {
                "date": target_date,
                "raw": payload,
                "source": "JiejiariAPI external API",
            }
        ]

    def generate_delay_risk_report(self, _: dict[str, Any]) -> list[dict[str, Any]]:
        now = now_text()
        risky = []
        for task in self.data_service.list_tasks():
            if task.current_status != "已完成" and (task.planned_end_at or "") < now:
                risky.append(task.model_dump())
        if not risky:
            risky = [task.model_dump() for task in self.data_service.list_tasks() if task.current_status != "已完成"][:5]
        return risky

    def _choose_tool_by_rules(self, question: str) -> str:
        q = question.upper()
        if any(key in question for key in ["天气", "下雨", "台风", "暴雨", "出货风险", "送货风险"]):
            return "query_weather"
        if self._looks_like_logistics_question(question):
            return "query_logistics"
        if any(key in question for key in ["路线", "距离", "几公里", "多久能到", "预计到达", "送货路线", "运输时间"]):
            return "query_route_plan"
        if any(key in question for key in ["节假日", "放假", "调休", "工作日", "周末", "法定假"]):
            return "query_workday_calendar"
        if any(key in question for key in ["联网", "搜索", "网上", "行业资料", "材料行情", "公开信息", "最新政策"]):
            return "query_web_search"
        if any(key in question for key in ["图片", "照片", "单据", "OCR", "识别", "缺陷图", "报警图", "标签", "图纸", "工艺单"]):
            return "query_multimodal_records"
        if any(key in question for key in ["库存", "ABS", "PP", "材料", "模具钢"]):
            return "query_inventory"
        if any(key in question for key in ["质检", "不良", "异常", "缩水", "毛边"]):
            return "query_quality_issues"
        if "注塑" in question or "注塑机" in question:
            return "query_injection_runs"
        if any(key in question for key in ["延期", "风险", "来不及"]):
            return "generate_delay_risk_report"
        if any(key in q for key in ["CNC", "火花", "线切割", "深孔", "车床", "铣床", "磨床", "摇臂", "设备", "排产"]):
            return "query_equipment_schedule"
        return "query_order_progress"

    @staticmethod
    def _force_tool_by_rules(question: str) -> str | None:
        if FactoryAgentService._looks_like_logistics_question(question):
            return "query_logistics"
        return None

    @staticmethod
    def _looks_like_logistics_question(question: str) -> bool:
        normalized = re.sub(r"\s+", "", question)
        if any(key in question for key in ["物流", "快递", "运单", "签收", "派送", "运输轨迹"]):
            return True
        if "单号" in question and re.search(r"[A-Za-z0-9]{8,30}", question):
            return True
        return bool(re.fullmatch(r"[A-Za-z0-9]{10,30}", normalized))

    def _format_answer(self, tool: str, data: list[dict[str, Any]]) -> str:
        if not data:
            return "没有查询到匹配数据。"
        if tool == "query_inventory":
            lines = ["库存情况："]
            for row in data:
                state = "低于安全库存" if float(row["qty"]) < float(row["safety_qty"]) else "正常"
                lines.append(f"- {row['material_name']} {row.get('spec') or ''}：{row['qty']}{row['unit']}，安全库存 {row['safety_qty']}{row['unit']}，状态：{state}。")
            return "\n".join(lines)
        if tool == "query_weather":
            if data and data[0].get("error"):
                return str(data[0]["error"])
            lines = ["外部天气查询结果："]
            for row in data:
                rain = row.get("rain_probability")
                advice = "注意防雨和物流延误" if rain is not None and rain >= 50 else "天气风险较低"
                lines.append(
                    f"- {row['city']} {row['date']}：{row['temp_min']}~{row['temp_max']}℃，降雨概率 {rain}% ，{advice}。"
                )
            return "\n".join(lines)
        if tool == "query_logistics":
            if data and data[0].get("error"):
                return str(data[0]["error"])
            lines = ["物流查询结果："]
            for row in data:
                if not row.get("success"):
                    lines.append(f"- 单号 {row['logistic_code']}：查询失败，原因：{row.get('reason') or '接口未返回原因'}。")
                    continue
                latest = row.get("latest_accept_station") or "暂无轨迹"
                lines.append(f"- 单号 {row['logistic_code']}：状态码 {row.get('state') or '未知'}，最新节点：{row.get('latest_accept_time') or '未知时间'} {latest}，轨迹数 {row.get('trace_count', 0)}。")
            return "\n".join(lines)
        if tool == "query_route_plan":
            if data and data[0].get("error"):
                return str(data[0]["error"])
            lines = ["路线规划结果："]
            for row in data:
                hours = round(float(row["duration_minutes"]) / 60, 1)
                lines.append(f"- {row['origin']} 到 {row['destination']}：约 {row['distance_km']} 公里，预计 {row['duration_minutes']} 分钟（约 {hours} 小时），策略：{row['strategy']}。")
            return "\n".join(lines)
        if tool == "query_web_search":
            if data and data[0].get("error"):
                return str(data[0]["error"])
            lines = ["联网搜索结果："]
            for row in data:
                if row.get("answer"):
                    lines.append(f"- 综合结论：{row['answer']}")
                for item in row.get("results", []):
                    lines.append(f"- {item['title']}：{item['content']} 来源：{item['url']}")
            return "\n".join(lines)
        if tool == "query_workday_calendar":
            if data and data[0].get("error"):
                return str(data[0]["error"])
            lines = ["节假日/工作日查询结果："]
            for row in data:
                summary = self._summarize_holiday_payload(row.get("raw") or {}, row["date"])
                lines.append(f"- {row['date']}：{summary}")
            return "\n".join(lines)
        if tool == "query_quality_issues":
            lines = ["质检异常："]
            for row in data:
                lines.append(f"- {row['product_name']}：{row['issue_type']}，不良 {row['bad_qty']}，状态 {row['status']}，原因：{row.get('cause') or '待分析'}。")
            return "\n".join(lines)
        if tool == "query_multimodal_records":
            lines = ["多模态识别记录："]
            for row in data:
                extracted = row.get("extracted") or {}
                fields = "，".join(f"{key}：{value}" for key, value in extracted.items() if value) or "暂无结构化字段"
                lines.append(f"- {row['filename']}（{MULTIMODAL_CATEGORIES.get(row['category'], row['category'])}）：{row['analysis']} 字段：{fields}。建议：{row['suggestions']}")
            return "\n".join(lines)
        if tool == "query_injection_runs":
            lines = ["注塑生产："]
            for row in data:
                lines.append(f"- {row['machine_code']} 生产 {row['product_name']}，材料 {row['material_name']}，计划 {row['plan_qty']}，已完成 {row['completed_qty']}，状态 {row['status']}。")
            return "\n".join(lines)
        if tool == "generate_delay_risk_report":
            lines = ["延期风险任务："]
            for row in data:
                lines.append(f"- {row['product_name']}：{row['process_name']}，设备 {row['equipment_code']}，状态 {row['current_status']}，计划结束 {row.get('planned_end_at') or '待确认'}。")
            return "\n".join(lines)
        if tool in {"query_process_tasks", "query_equipment_schedule"}:
            lines = ["工序/设备进度："]
            for row in data:
                lines.append(f"- {row['product_name']}：{row['process_name']}，设备 {row['equipment_code']}，状态 {row['current_status']}，负责人 {row['owner_name']}。")
            return "\n".join(lines)
        lines = ["订单/模具进度："]
        for row in data:
            if "order_no" in row:
                lines.append(f"- 订单 {row['order_no']}：{row['product_name']}，数量 {row['quantity']}，交期 {row['due_date']}，状态 {row['status']}。")
            elif "mold_no" in row:
                lines.append(f"- 模具 {row['mold_no']}：{row['mold_name']}，当前阶段 {row['current_stage']}，状态 {row['status']}。")
        return "\n".join(lines)

    def _log_tool(self, question: str, tool_name: str, arguments: dict[str, Any]) -> None:
        with self.data_service.database.connect() as conn:
            conn.execute(
                "INSERT INTO agent_tool_logs VALUES (?, ?, ?, ?, ?)",
                (uuid.uuid4().hex, question, tool_name, json.dumps(arguments, ensure_ascii=False), now_text()),
            )

    def _filter_by_context(self, rows: list[dict[str, Any]], args: dict[str, Any]) -> list[dict[str, Any]]:
        question = str(args.get("question", ""))
        context = self._normalize_context(args.get("context") or {})
        subject = self._find_subject(question) or context
        if not subject:
            return rows
        order_no = str(subject.get("order_no") or "")
        product_name = str(subject.get("product_name") or "")
        mold_no = str(subject.get("mold_no") or "")
        mold_name = str(subject.get("mold_name") or "")

        def row_matches(row: dict[str, Any]) -> bool:
            text = " ".join(str(value or "") for value in row.values())
            if order_no and (row.get("order_no") == order_no or row.get("order_id") == subject.get("order_id") or order_no in text):
                return True
            if mold_no and mold_no in text:
                return True
            for name in [product_name, mold_name]:
                if name and (name in text or text in name):
                    return True
                if name.endswith("模具") and name[:-2] and name[:-2] in text:
                    return True
                if name.endswith("塑胶件") and name[:-3] and name[:-3] in text:
                    return True
            return False

        filtered = [row for row in rows if row_matches(row)]
        return filtered or rows

    def _find_subject(self, question: str) -> dict[str, Any]:
        q = question.upper()
        for order in self.data_service.list_orders():
            if order.order_no.upper() in q or order.product_name in question:
                return {
                    "order_id": order.id,
                    "order_no": order.order_no,
                    "product_name": order.product_name,
                }
        for mold in self.data_service.list_molds():
            base_name = mold.mold_name.replace("模具", "")
            if mold.mold_no.upper() in q or mold.mold_name in question or (base_name and base_name in question):
                return {
                    "order_id": mold.order_id,
                    "mold_no": mold.mold_no,
                    "mold_name": mold.mold_name,
                    "product_name": base_name or mold.mold_name,
                }
        for task in self.data_service.list_tasks():
            product_name = task.product_name.replace("模具", "").replace("塑胶件", "")
            if task.product_name in question or (product_name and product_name in question):
                return {
                    "order_id": task.order_id,
                    "mold_id": task.mold_id,
                    "product_name": product_name or task.product_name,
                }
        return {}

    def _context_from_result(self, question: str, data: list[dict[str, Any]], previous: dict[str, Any]) -> dict[str, Any]:
        subject = self._find_subject(question)
        if subject:
            return self._normalize_context({**previous, **subject})
        candidates = []
        for row in data:
            if row.get("order_no") or row.get("mold_no") or row.get("product_name") or row.get("mold_name"):
                candidates.append(row)
        if len(candidates) == 1:
            row = candidates[0]
            next_context = {**previous}
            for key in ["order_id", "order_no", "mold_id", "mold_no", "product_name", "mold_name"]:
                if row.get(key):
                    next_context[key] = row[key]
            return self._normalize_context(next_context)
        return self._normalize_context(previous)

    @staticmethod
    def _normalize_context(context: dict[str, Any]) -> dict[str, Any]:
        allowed = {"order_id", "order_no", "mold_id", "mold_no", "product_name", "mold_name"}
        return {key: value for key, value in context.items() if key in allowed and value}

    def _with_context(self, question: str, context: dict[str, Any]) -> str:
        if not context:
            return question
        if self._find_subject(question):
            return question
        subject_parts = []
        for label, key in [("订单", "order_no"), ("产品", "product_name"), ("模具", "mold_name"), ("模具编号", "mold_no")]:
            if context.get(key):
                subject_parts.append(f"{label}{context[key]}")
        if not subject_parts:
            return question
        return f"{question}（当前上下文：{'，'.join(subject_parts)}）"

    @staticmethod
    def _extract_logistic_code(question: str) -> str:
        patterns = [
            r"(?:物流单号|快递单号|运单号|单号)[:：]?\s*([A-Za-z0-9]{8,30})",
            r"\b([A-Za-z0-9]{10,30})\b",
            r"([A-Za-z]*\d[A-Za-z0-9]{7,29})",
        ]
        for pattern in patterns:
            match = re.search(pattern, question)
            if match:
                return match.group(1).strip()
        return ""

    @staticmethod
    def _guess_shipper_code(question: str) -> str:
        upper = question.upper()
        for name, code in SHIPPER_CODES.items():
            if name.upper() in upper or code in upper:
                return code
        return ""

    @staticmethod
    def _extract_route_points(question: str) -> tuple[str, str]:
        patterns = [
            r"(?:从)?(?P<origin>[\u4e00-\u9fa5A-Za-z0-9\s]{2,30}?)(?:发到|送到|运到|到|去|发)(?P<destination>[\u4e00-\u9fa5A-Za-z0-9\s]{2,30}?)(?:的?送货路线|的?路线|送货.*|要多久|多久.*|距离.*|几公里.*|多少公里.*|多远.*|路程.*|运输时间.*|$)",
            r"(?P<origin>[\u4e00-\u9fa5A-Za-z0-9\s]+?)(?:发到|送到|运到)(?P<destination>[\u4e00-\u9fa5A-Za-z0-9\s]+)",
        ]
        for pattern in patterns:
            match = re.search(pattern, question)
            if match:
                origin = FactoryAgentService._clean_place_candidate(match.group("origin"))
                destination = FactoryAgentService._clean_place_candidate(match.group("destination"))
                if origin and destination and origin != destination:
                    return origin, destination
        destination = FactoryAgentService._guess_city(question)
        origin = os.getenv("DEFAULT_FACTORY_ADDRESS", "深圳")
        if destination and destination != origin and any(key in question for key in ["发到", "送到", "运到", "送货", "路线"]):
            return origin, destination
        return "", ""

    @staticmethod
    def _clean_place_candidate(value: str) -> str:
        place = value.strip(" ，。！？?：:")
        for word in [
            "查询", "查一下", "看看", "看下", "工厂", "客户", "这批货", "货物", "产品", "订单",
            "送货路线", "路线", "送货", "发货", "大概", "大约", "大概要", "需要", "要", "多久",
            "多久能到", "距离", "几公里", "多少公里", "多远", "路程", "运输时间", "的",
        ]:
            place = place.replace(word, "")
        return place.strip(" ，。！？?：:")

    @staticmethod
    def _is_ambiguous_province(place: str) -> bool:
        normalized = place.removesuffix("省").removesuffix("市").removesuffix("自治区").removesuffix("特别行政区")
        return normalized in PROVINCE_NAMES

    def _amap_geocode(self, address: str, api_key: str) -> dict[str, Any]:
        url = "https://restapi.amap.com/v3/geocode/geo?" + urlencode({"key": api_key, "address": address})
        payload = self._load_external_json(url)
        if payload.get("status") != "1":
            return {"error": payload.get("info") or f"高德无法解析地址：{address}"}
        geocodes = payload.get("geocodes") or []
        if not geocodes:
            return {"error": f"高德无法解析地址：{address}"}
        return {"formatted_address": geocodes[0].get("formatted_address") or address, "location": geocodes[0]["location"]}

    @staticmethod
    def _extract_date(question: str) -> str:
        now = datetime.now()
        if "后天" in question:
            return (now + timedelta(days=2)).strftime("%Y-%m-%d")
        if "明天" in question:
            return (now + timedelta(days=1)).strftime("%Y-%m-%d")
        if "今天" in question:
            return now.strftime("%Y-%m-%d")
        match = re.search(r"(\d{4})[-/年](\d{1,2})[-/月](\d{1,2})日?", question)
        if match:
            year, month, day = match.groups()
            return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"
        match = re.search(r"(\d{1,2})月(\d{1,2})日", question)
        if match:
            month, day = match.groups()
            return f"{now.year:04d}-{int(month):02d}-{int(day):02d}"
        return now.strftime("%Y-%m-%d")

    @staticmethod
    def _summarize_holiday_payload(payload: dict[str, Any], date_text: str) -> str:
        data = payload.get("data", payload)
        is_weekend = False
        try:
            is_weekend = datetime.strptime(date_text, "%Y-%m-%d").weekday() >= 5
        except ValueError:
            pass
        if isinstance(data, dict):
            name = data.get("name") or data.get("holiday_name") or data.get("holidayName") or data.get("type") or ""
            is_holiday = data.get("is_holiday", data.get("isHoliday", data.get("holiday")))
            is_workday = data.get("is_workday", data.get("isWorkday", data.get("workday")))
            if is_holiday is True or str(is_holiday).lower() in {"true", "1", "yes"}:
                return f"节假日{name and '，' + str(name)}，排产和交付需要预留缓冲。"
            if is_workday is True or str(is_workday).lower() in {"true", "1", "yes"}:
                return "工作日，可按正常生产和物流计划评估。"
            if is_weekend:
                return "周末，接口未标记为法定节假日，但交付和物流仍建议预留缓冲。"
            if is_holiday is False or str(is_holiday).lower() in {"false", "0", "no"}:
                return "非法定节假日，按普通工作日风险评估；如企业周末休息需单独考虑。"
        return f"接口返回：{json.dumps(payload, ensure_ascii=False)}"

    @staticmethod
    def _guess_city(question: str) -> str:
        known_cities = [
            "北京", "上海", "天津", "重庆", "深圳", "广州", "东莞", "佛山", "惠州", "中山", "珠海",
            "杭州", "宁波", "温州", "苏州", "昆山", "无锡", "南京", "合肥", "福州", "厦门",
            "长沙", "武汉", "南昌", "郑州", "成都", "绵阳", "西安", "青岛", "济南", "烟台",
            "沈阳", "大连", "长春", "哈尔滨", "石家庄", "太原", "呼和浩特", "南宁", "海口",
            "贵阳", "昆明", "兰州", "银川", "西宁", "乌鲁木齐", "拉萨",
        ]
        for city in known_cities:
            if city in question:
                return city
        patterns = [
            r"(?:查一下|查询|看看|看下)?(?P<city>[\u4e00-\u9fa5]{2,10})(?:今天|明天|后天|这几天|未来几天|本周|周末)?(?:天气|气温|温度)",
            r"(?P<city>[\u4e00-\u9fa5]{2,10})(?:今天|明天|后天|这几天|未来几天|本周|周末)?(?:会不会|有没有|是否)?(?:下雨|降雨|暴雨|台风|下雪)",
            r"(?:发往|发到|送到|运到|去|到)(?P<city>[\u4e00-\u9fa5]{2,10})(?:.*?)(?:天气|下雨|降雨|暴雨|台风|出货|送货)",
        ]
        for pattern in patterns:
            match = re.search(pattern, question)
            if match:
                city = FactoryAgentService._clean_city_candidate(match.group("city"))
                if city:
                    return city
        return os.getenv("DEFAULT_WEATHER_CITY", "深圳")

    @staticmethod
    def _clean_city_candidate(value: str) -> str:
        city = value.strip(" ，。！？?：:")
        for word in ["今天", "明天", "后天", "未来几天", "这几天", "本周", "周末", "现在", "当地", "那边", "那里", "查一下", "查询", "看看", "看下"]:
            city = city.replace(word, "")
        for suffix in ["市", "地区"]:
            if city.endswith(suffix) and len(city) > len(suffix) + 1:
                city = city[: -len(suffix)]
        return city.strip(" ，。！？?：:")

    @staticmethod
    def _load_external_json(url: str) -> Any:
        request = urllib.request.Request(
            url,
            headers={
                "User-Agent": "factory-agent-project/1.0",
                "Accept": "application/json",
            },
        )
        try:
            with urllib.request.urlopen(request, timeout=10) as response:
                return json.loads(response.read().decode("utf-8"))
        except Exception:
            logger.exception("external_json_request_failed url=%s", FactoryAgentService._redact_secret_url(url))
            raise

    @staticmethod
    def _post_json(url: str, payload: dict[str, Any], headers: dict[str, str] | None = None) -> Any:
        request_headers = {
            "User-Agent": "factory-agent-project/1.0",
            "Accept": "application/json",
            "Content-Type": "application/json",
        }
        request_headers.update(headers or {})
        request = urllib.request.Request(
            url,
            data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            headers=request_headers,
            method="POST",
        )
        try:
            with urllib.request.urlopen(request, timeout=15) as response:
                return json.loads(response.read().decode("utf-8"))
        except Exception:
            logger.exception("external_json_post_failed url=%s", url)
            raise

    @staticmethod
    def _post_form_json(url: str, form: dict[str, Any]) -> Any:
        request = urllib.request.Request(
            url,
            data=urlencode(form).encode("utf-8"),
            headers={
                "User-Agent": "factory-agent-project/1.0",
                "Accept": "application/json",
                "Content-Type": "application/x-www-form-urlencoded;charset=utf-8",
            },
            method="POST",
        )
        try:
            with urllib.request.urlopen(request, timeout=15) as response:
                return json.loads(response.read().decode("utf-8"))
        except Exception:
            logger.exception("external_form_post_failed url=%s", url)
            raise

    @staticmethod
    def _redact_secret_url(url: str) -> str:
        return re.sub(r"([?&](?:key|api_key|apikey|sig|signature)=)[^&]+", r"\1***", url, flags=re.I)
